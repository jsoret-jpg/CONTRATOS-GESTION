from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="333333", end_color="333333", fill_type="solid")


def _write_header(ws, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col_idx)].width = max(18, len(header) + 4)


def build_template_workbook(headers: list[str], example_row: list, instructions: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    _write_header(ws, headers)

    for col_idx, value in enumerate(example_row, start=1):
        ws.cell(row=2, column=col_idx, value=value)

    instructions_ws = wb.create_sheet("Instrucciones")
    instructions_ws.column_dimensions["A"].width = 100
    for row_idx, text in enumerate(instructions, start=1):
        instructions_ws.cell(row=row_idx, column=1, value=text)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_data_workbook(headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    _write_header(ws, headers)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
